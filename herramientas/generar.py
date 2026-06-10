#!/usr/bin/env python3
"""Regenera index.html a partir del Excel maestro (Programación MP) y plantilla.html.

Uso:
    pip install openpyxl
    python3 herramientas/generar.py ProgramacionMP_2026.xlsm

Lee las hojas PMP_2026 y Registro_MP-2026 e inyecta los datos en la
aplicación. Sirve para partir un año nuevo o re-sincronizar con el maestro
(ojo: regenerar y reemplazar los datos borra lo registrado en la app si
luego se usa "Restablecer"; exportar respaldo antes).
"""
import json
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
HOJA_PMP = "PMP_2026"
HOJA_REG = "Registro_MP-2026"


def limpio(v):
    if v is None:
        return ""
    v = str(v).strip()
    return "" if v in ("0", "N/A", "n/a") else v


def extraer(ruta_xlsm):
    wb = openpyxl.load_workbook(ruta_xlsm, read_only=True, data_only=True)
    ws = wb[HOJA_PMP]
    equipos = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if row[1] is None:
            continue
        try:
            rid = int(row[1])
        except (TypeError, ValueError):
            continue
        equipos.append({
            "id": rid, "carpeta": limpio(row[2]), "inv": limpio(row[3]),
            "fam": limpio(row[0]), "equipo": limpio(row[4]),
            "servicio": limpio(row[5]), "unidad": limpio(row[6]),
            "ubicacion": limpio(row[7]), "proc": limpio(row[8]),
            "marca": limpio(row[9]), "modelo": limpio(row[10]),
            "serie": str(row[11]).strip() if row[11] is not None else "",
            "anio": limpio(row[12]), "clasif": limpio(row[14]),
            "obs": limpio(row[16]), "frec": limpio(row[17]),
            "respMP": limpio(row[18]),
            "plan": [limpio(row[19 + i]).upper() for i in range(12)],
        })
    # la columna Familia viene con celdas combinadas: arrastrar hacia abajo
    ultima = ""
    for e in equipos:
        if e["fam"]:
            ultima = e["fam"]
        else:
            e["fam"] = ultima

    ws2 = wb[HOJA_REG]
    registro = {}
    for row in ws2.iter_rows(min_row=8, values_only=True):
        if row[1] is None:
            continue
        try:
            rid = int(row[1])
        except (TypeError, ValueError):
            continue
        registro[rid] = {
            "res": [[limpio(row[19 + 2 * i]).upper(), limpio(row[20 + 2 * i])]
                    for i in range(12)],
            "obs": limpio(row[43]),
        }
    for e in equipos:
        r = registro.get(e["id"])
        e["resultado"] = r["res"] if r else [["", ""]] * 12
        e["obsReg"] = r["obs"] if r else ""
    return equipos


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    equipos = extraer(sys.argv[1])
    datos = json.dumps(equipos, ensure_ascii=False,
                       separators=(",", ":")).replace("</", "<\\/")
    plantilla = (RAIZ / "plantilla.html").read_text(encoding="utf-8")
    (RAIZ / "index.html").write_text(
        plantilla.replace("__DATA__", datos), encoding="utf-8")
    print(f"index.html generado con {len(equipos)} equipos.")


if __name__ == "__main__":
    main()
